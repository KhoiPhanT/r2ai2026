from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from legal_rag.corpus.ingest import ingest_corpus
from legal_rag.documents.filenames import normalized_docx_name, slug_doc_id, slug_title
from legal_rag.documents.normalizer import normalize_documents


class DocumentNormalizerTest(unittest.TestCase):
    def test_slug_helpers(self) -> None:
        self.assertEqual(slug_doc_id("40/2024/QH15"), "40-2024-QH15")
        self.assertNotIn("/", normalized_docx_name("40/2024/QH15", "Sửa đổi, bổ sung Luật Cảnh vệ"))
        self.assertLessEqual(len(normalized_docx_name("40/2024/QH15", "a" * 300)), 165)
        self.assertEqual(slug_title("Luật Cảnh vệ"), "luat-canh-ve")
        self.assertEqual(slug_title("Sửa đổi Luật Thủ đô"), "sua-doi-luat-thu-do")

    def test_normalize_docx_and_ingest_jsonl(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source_dir = root / "raw"
            source_dir.mkdir()
            docx_path = source_dir / "40_2024_QH15_564046.docx"
            self._write_sample_docx(docx_path)

            output_dir = root / "normalized"
            result = normalize_documents(source_dir, output_dir)

            self.assertEqual(len(result.documents), 1)
            record = result.documents[0]
            self.assertEqual(record.doc_id, "40/2024/QH15")
            self.assertEqual(record.doc_type, "Luật")
            self.assertEqual(record.trich_yeu, "SỬA ĐỔI, BỔ SUNG MỘT SỐ ĐIỀU CỦA LUẬT CẢNH VỆ")
            self.assertIn("40/2024/QH15", record.title_for_submission)
            self.assertTrue((output_dir / "documents.jsonl").exists())
            self.assertTrue((output_dir / "manifest.json").exists())
            self.assertTrue(any((output_dir / "docx").glob("40-2024-QH15__*.docx")))

            articles, warnings = ingest_corpus(output_dir / "documents.jsonl")
            self.assertEqual(warnings, [])
            self.assertEqual(len(articles), 2)

    def test_missing_required_metadata_is_reported_as_failure(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source_dir = root / "raw"
            source_dir.mkdir()
            docx_path = source_dir / "broken.docx"
            document = Document()
            document.add_paragraph("Một văn bản không có title block hợp lệ.")
            document.save(docx_path)

            result = normalize_documents(source_dir, root / "normalized")

            self.assertEqual(result.documents, [])
            self.assertEqual(len(result.report["failures"]), 1)
            self.assertIn("missing_required_field:doc_id", result.report["failures"][0]["error"])

    def test_normalize_bo_luat_title_block(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source_dir = root / "raw"
            source_dir.mkdir()
            docx_path = source_dir / "45_2019_QH14_333670.docx"

            document = Document()
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "QUỐC HỘI"
            table.cell(0, 1).text = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"
            table.cell(1, 0).text = "Luật số: 45/2019/QH14"
            table.cell(1, 1).text = "Hà Nội, ngày 20 tháng 11 năm 2019"
            document.add_paragraph("BỘ LUẬT")
            document.add_paragraph("LAO ĐỘNG")
            document.add_paragraph("Căn cứ Hiến pháp nước Cộng hòa xã hội chủ nghĩa Việt Nam;")
            document.add_paragraph("Điều 1. Phạm vi điều chỉnh")
            document.add_paragraph("Bộ luật Lao động quy định tiêu chuẩn lao động.")
            document.save(docx_path)

            result = normalize_documents(source_dir, root / "normalized")

            self.assertEqual(len(result.documents), 1)
            record = result.documents[0]
            self.assertEqual(record.doc_id, "45/2019/QH14")
            self.assertEqual(record.doc_type, "Bộ luật")
            self.assertEqual(record.trich_yeu, "LAO ĐỘNG")
            self.assertEqual(record.title_for_submission, "Bộ luật 45/2019/QH14 LAO ĐỘNG")

    @staticmethod
    def _write_sample_docx(path: Path) -> None:
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "QUỐC HỘI"
        table.cell(0, 1).text = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"
        table.cell(1, 0).text = "Luật số: 40/2024/QH15"
        table.cell(1, 1).text = "Hà Nội, ngày 28 tháng 6 năm 2024"
        document.add_paragraph("LUẬT")
        document.add_paragraph("SỬA ĐỔI, BỔ SUNG MỘT SỐ ĐIỀU CỦA LUẬT CẢNH VỆ")
        document.add_paragraph("Căn cứ Hiến pháp nước Cộng hòa xã hội chủ nghĩa Việt Nam;")
        document.add_paragraph("Điều 1. Phạm vi điều chỉnh")
        document.add_paragraph("Luật này sửa đổi một số quy định.")
        document.add_paragraph("Điều 2. Hiệu lực thi hành")
        document.add_paragraph("Luật này có hiệu lực từ ngày 01 tháng 01 năm 2025.")
        document.save(path)


if __name__ == "__main__":
    unittest.main()
