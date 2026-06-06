from __future__ import annotations

import json
import re
import shutil
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from legal_rag.documents.filenames import normalized_docx_name
from legal_rag.schemas.models import DocumentRecord
from legal_rag.utils.text import extract_article_labels, normalize_text

DOC_ID_PATTERN = re.compile(r"\b(\d{1,4})[/_](\d{4})[/_]([A-ZĐ0-9]+)\b", re.IGNORECASE)
FILENAME_DOC_ID_PATTERN = re.compile(r"(\d{1,4})_(\d{4})_([A-ZĐ0-9]+)", re.IGNORECASE)
DOC_TYPE_CANDIDATES = (
    "BỘ LUẬT",
    "LUẬT",
    "NGHỊ QUYẾT",
    "NGHỊ ĐỊNH",
    "THÔNG TƯ",
    "QUYẾT ĐỊNH",
    "PHÁP LỆNH",
)
TITLE_STOP_PREFIXES = (
    "Căn cứ",
    "Quốc hội ban hành",
    "Chính phủ ban hành",
    "Chương ",
    "Điều ",
    "Article ",
    "Pursuant ",
)


@dataclass(slots=True)
class NormalizeResult:
    documents: list[DocumentRecord]
    manifest: list[dict[str, Any]]
    report: dict[str, Any]


@dataclass(slots=True)
class ParsedDocx:
    source_path: Path
    doc_id: str
    doc_type: str
    trich_yeu: str
    title_for_submission: str
    raw_text: str
    article_count: int
    paragraphs: list[str]
    header_text: str
    warnings: list[str]


def normalize_documents(input_path: str | Path, output_dir: str | Path) -> NormalizeResult:
    input_root = Path(input_path)
    output_root = Path(output_dir)
    docx_dir = output_root / "docx"
    txt_dir = output_root / "txt"
    reports_dir = output_root / "reports"
    for path in (docx_dir, txt_dir, reports_dir):
        if path.exists():
            shutil.rmtree(path)
        path.mkdir(parents=True, exist_ok=True)

    source_files = _iter_document_files(input_root)
    parsed_docx: list[ParsedDocx] = []
    legacy_docs: list[dict[str, Any]] = []
    failures: list[dict[str, Any]] = []

    for source in source_files:
        if source.suffix.lower() == ".docx":
            try:
                parsed_docx.append(_parse_docx(source))
            except Exception as exc:  # noqa: BLE001
                failures.append({"source_path": str(source), "error": f"{type(exc).__name__}: {exc}"})
        elif source.suffix.lower() == ".doc":
            legacy_docs.append(_process_legacy_doc(source, txt_dir))

    doc_id_counts = Counter(item.doc_id for item in parsed_docx)
    seen_doc_ids: defaultdict[str, int] = defaultdict(int)
    documents: list[DocumentRecord] = []
    manifest: list[dict[str, Any]] = []
    duplicate_doc_ids = sorted(doc_id for doc_id, count in doc_id_counts.items() if count > 1)

    for parsed in parsed_docx:
        duplicate_index = seen_doc_ids[parsed.doc_id]
        seen_doc_ids[parsed.doc_id] += 1
        warnings = list(parsed.warnings)
        if doc_id_counts[parsed.doc_id] > 1:
            warnings.append(f"duplicate_doc_id:{parsed.doc_id}")

        file_name = normalized_docx_name(parsed.doc_id, parsed.trich_yeu, duplicate_index)
        normalized_path = docx_dir / file_name
        shutil.copy2(parsed.source_path, normalized_path)
        txt_path = txt_dir / f"{normalized_path.stem}.txt"
        txt_path.write_text(parsed.raw_text, encoding="utf-8")

        record = DocumentRecord(
            doc_id=parsed.doc_id,
            doc_type=_title_case_doc_type(parsed.doc_type),
            trich_yeu=parsed.trich_yeu,
            title_for_submission=parsed.title_for_submission,
            raw_text=parsed.raw_text,
            source_url=str(parsed.source_path),
            metadata={
                "source_path": str(parsed.source_path),
                "normalized_docx_path": str(normalized_path),
                "normalized_txt_path": str(txt_path),
                "article_count": parsed.article_count,
                "header_text": parsed.header_text,
            },
        )
        documents.append(record)
        manifest.append(
            {
                "source_path": str(parsed.source_path),
                "normalized_docx_path": str(normalized_path),
                "normalized_txt_path": str(txt_path),
                "doc_id": record.doc_id,
                "doc_type": record.doc_type,
                "trich_yeu": record.trich_yeu,
                "title_for_submission": record.title_for_submission,
                "article_count": parsed.article_count,
                "warnings": warnings,
            }
        )

    report = {
        "input_path": str(input_root),
        "output_dir": str(output_root),
        "source_files": len(source_files),
        "docx_files": len([p for p in source_files if p.suffix.lower() == ".docx"]),
        "legacy_doc_files": len([p for p in source_files if p.suffix.lower() == ".doc"]),
        "documents_written": len(documents),
        "failures": failures,
        "legacy_docs": legacy_docs,
        "duplicate_doc_ids": duplicate_doc_ids,
        "warnings_count": sum(len(item.get("warnings", [])) for item in manifest)
        + sum(len(item.get("warnings", [])) for item in legacy_docs),
    }
    _write_outputs(output_root, reports_dir, documents, manifest, report)
    return NormalizeResult(documents=documents, manifest=manifest, report=report)


def _parse_docx(path: Path) -> ParsedDocx:
    doc = Document(str(path))
    blocks = list(_iter_block_text(doc))
    all_text = normalize_text("\n\n".join(text for text in blocks if text.strip()))
    paragraphs = [normalize_text(p.text) for p in doc.paragraphs if normalize_text(p.text)]
    header_cells = _first_table_cells(doc)
    header_text = normalize_text("\n".join(header_cells + paragraphs[:8]))
    doc_id = _extract_doc_id(header_text) or _extract_doc_id_from_filename(path.name)
    doc_type, title_lines = _extract_type_and_title(paragraphs)
    trich_yeu = normalize_text(" ".join(title_lines))
    warnings: list[str] = []

    if not doc_id:
        warnings.append("missing_required_field:doc_id")
    if not doc_type:
        warnings.append("missing_required_field:doc_type")
    if not trich_yeu:
        warnings.append("missing_required_field:trich_yeu")
    if not all_text:
        warnings.append("missing_required_field:raw_text")
    required_errors = [warning for warning in warnings if warning.startswith("missing_required_field:")]
    if required_errors:
        raise ValueError(", ".join(required_errors))
    article_count = len(extract_article_labels(all_text))
    if article_count == 0:
        warnings.append("no_article_labels")

    doc_type_title = _title_case_doc_type(doc_type)
    title_for_submission = normalize_text(f"{doc_type_title} {doc_id} {trich_yeu}")
    if doc_id not in title_for_submission:
        warnings.append(f"title_missing_doc_id:{doc_id}")

    return ParsedDocx(
        source_path=path,
        doc_id=doc_id,
        doc_type=doc_type,
        trich_yeu=trich_yeu,
        title_for_submission=title_for_submission,
        raw_text=all_text,
        article_count=article_count,
        paragraphs=paragraphs,
        header_text=header_text,
        warnings=warnings,
    )


def _process_legacy_doc(path: Path, txt_dir: Path) -> dict[str, Any]:
    txt_path = txt_dir / f"{path.stem}.legacy.txt"
    record: dict[str, Any] = {
        "source_path": str(path),
        "normalized_txt_path": str(txt_path),
        "status": "excluded_legacy_doc",
        "warnings": ["legacy_doc_excluded"],
    }
    try:
        import subprocess

        completed = subprocess.run(
            ["textutil", "-convert", "txt", "-stdout", str(path)],
            check=True,
            capture_output=True,
            text=True,
        )
        text = normalize_text(completed.stdout)
        txt_path.write_text(text, encoding="utf-8")
        if _looks_english(text):
            record["warnings"].append("language_english_excluded")
        record["chars"] = len(text)
    except Exception as exc:  # noqa: BLE001
        record["status"] = "failed_legacy_doc"
        record["error"] = f"{type(exc).__name__}: {exc}"
    return record


def _iter_document_files(input_root: Path) -> list[Path]:
    if input_root.is_file():
        return [input_root]
    return sorted(
        path
        for path in input_root.rglob("*")
        if path.is_file() and path.suffix.lower() in {".docx", ".doc"} and not path.name.startswith("~$")
    )


def _iter_block_text(document: DocumentObject) -> Iterable[str]:
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            text = normalize_text(Paragraph(child, document).text)
            if text:
                yield text
        elif isinstance(child, CT_Tbl):
            table = Table(child, document)
            rows = []
            for row in table.rows:
                cells = [normalize_text(cell.text).replace("\n", " ") for cell in row.cells]
                row_text = "\t".join(cell for cell in cells if cell)
                if row_text:
                    rows.append(row_text)
            if rows:
                yield "\n".join(rows)


def _first_table_cells(doc: DocumentObject) -> list[str]:
    cells: list[str] = []
    if not doc.tables:
        return cells
    for row in doc.tables[0].rows[:4]:
        for cell in row.cells:
            text = normalize_text(cell.text).replace("\n", " ")
            if text:
                cells.append(text)
    return cells


def _extract_doc_id(text: str) -> str:
    match = DOC_ID_PATTERN.search(text.replace("_", "/"))
    if not match:
        return ""
    return f"{match.group(1)}/{match.group(2)}/{match.group(3).upper()}"


def _extract_doc_id_from_filename(name: str) -> str:
    match = FILENAME_DOC_ID_PATTERN.search(name)
    if not match:
        return ""
    return f"{match.group(1)}/{match.group(2)}/{match.group(3).upper()}"


def _extract_type_and_title(paragraphs: list[str]) -> tuple[str, list[str]]:
    doc_type = ""
    title_lines: list[str] = []
    start_index = -1
    for index, paragraph in enumerate(paragraphs[:20]):
        upper = paragraph.upper().strip()
        for candidate in DOC_TYPE_CANDIDATES:
            if upper == candidate:
                doc_type = candidate
                start_index = index + 1
                break
            if upper.startswith(candidate + " "):
                doc_type = candidate
                remainder = paragraph[len(candidate) :].strip(" :-\t")
                if remainder:
                    title_lines.append(remainder)
                start_index = index + 1
                break
        if doc_type:
            break

    if not doc_type:
        return "", []

    for paragraph in paragraphs[start_index : start_index + 8]:
        if _is_title_stop(paragraph):
            break
        cleaned = normalize_text(paragraph)
        if cleaned and cleaned != "\xa0":
            title_lines.append(cleaned)
    return doc_type, title_lines


def _is_title_stop(text: str) -> bool:
    stripped = text.strip()
    return any(stripped.startswith(prefix) for prefix in TITLE_STOP_PREFIXES)


def _title_case_doc_type(doc_type: str) -> str:
    mapping = {
        "BỘ LUẬT": "Bộ luật",
        "LUẬT": "Luật",
        "NGHỊ QUYẾT": "Nghị quyết",
        "NGHỊ ĐỊNH": "Nghị định",
        "THÔNG TƯ": "Thông tư",
        "QUYẾT ĐỊNH": "Quyết định",
        "PHÁP LỆNH": "Pháp lệnh",
        "VĂN BẢN": "Văn bản",
    }
    return mapping.get(doc_type.upper(), doc_type.capitalize())


def _looks_english(text: str) -> bool:
    sample = text[:4000].lower()
    english_hits = sum(token in sample for token in ["the national assembly", "law no.", "article ", "pursuant to"])
    vietnamese_hits = sum(token in sample for token in ["quốc hội", "luật số", "điều ", "căn cứ"])
    return english_hits > vietnamese_hits


def _write_outputs(
    output_root: Path,
    reports_dir: Path,
    documents: list[DocumentRecord],
    manifest: list[dict[str, Any]],
    report: dict[str, Any],
) -> None:
    documents_path = output_root / "documents.jsonl"
    with documents_path.open("w", encoding="utf-8") as f:
        for doc in documents:
            payload = asdict(doc)
            payload.update(doc.metadata)
            payload.pop("metadata", None)
            f.write(json.dumps(payload, ensure_ascii=False) + "\n")

    (output_root / "manifest.json").write_text(
        json.dumps({"documents": manifest}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (reports_dir / "normalize_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
